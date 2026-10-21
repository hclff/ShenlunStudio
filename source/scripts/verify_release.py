"""Exercise the actual frozen EXE twice without network or model credentials."""
import hashlib
import argparse
import json
import os
import shutil
from pathlib import Path
import sqlite3
import subprocess
import sys
import time
from datetime import datetime
import zipfile
from xml.etree import ElementTree

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from shenlun.storage import Store
from shenlun import __version__


def intake_fixtures(work):
    """Real Word, PNG, mixed text/scan PDF and question JSON; no stubbed readers."""
    from docx import Document
    from PySide6.QtCore import Qt
    from PySide6.QtGui import QFont,QImage,QPainter,QPdfWriter
    from PySide6.QtWidgets import QApplication
    from shenlun.ui.widgets import ensure_fonts
    app=QApplication.instance() or QApplication([])
    ensure_fonts()
    image=QImage(1600,280,QImage.Format.Format_RGB32);image.fill(Qt.GlobalColor.white)
    painter=QPainter(image);painter.setPen(Qt.GlobalColor.black);painter.setFont(QFont('Microsoft YaHei',46))
    painter.drawText(50,145,'给定材料：社区保留电话服务。');painter.end()
    png=work/"截图'a&b.png"
    if not image.save(str(png)):raise RuntimeError('测试图片生成失败')
    pdf=work/'文字与扫描.pdf';writer=QPdfWriter(str(pdf));writer.setResolution(100)
    painter=QPainter(writer);painter.setFont(QFont('Microsoft YaHei',36))
    painter.drawText(60,120,'第一材料：保留人工窗口，服务老人。');writer.newPage()
    painter.drawImage(0,80,image.scaledToWidth(700));painter.end()
    word=work/'题目表格.docx';document=Document();document.add_paragraph('题干：概括社区做法。')
    table=document.add_table(rows=1,cols=2);table.cell(0,0).text='对象';table.cell(0,1).text='老年群众'
    document.save(word)
    structured=work/'结构化题目.json'
    structured.write_text(json.dumps({'question':'概括便民服务做法','materials':'保留人工窗口和电话服务','max_score':20},ensure_ascii=False),encoding='utf8')
    from PIL import Image
    with Image.open(png) as page:
        webp=work/'截图.webp';page.save(webp,'WEBP')
        tiff=work/'两页扫描.tiff';page.save(tiff,'TIFF',save_all=True,append_images=[page.copy()])
    return [word,pdf,png,structured,webp,tiff]


def verify_normal_launch(exe, folder, work, env):
    """Launch the frozen main EXE with no arguments, redirecting only the test AppData."""
    local = work / 'normal-launch-localappdata'
    target = local / 'ShenlunStudio' / 'studio.db'
    store = Store(target)
    store.set_setting('normal_launch_preservation', {'value':'保留正常入口的设置'})
    store.upsert_article({'id':'normal-user-article', 'title':'自有文章', 'body':'原有正文保持不变。'})
    review = store.save_review({'submission':{'answer':'正常入口原答'}, 'assessment':{'total_score':None}})
    config = {**env, 'LOCALAPPDATA':str(local), 'QT_QPA_PLATFORM':'offscreen'}
    for attempt in range(2):
        process = subprocess.Popen([str(exe)], cwd=folder, env=config,
                                   creationflags=getattr(subprocess,'CREATE_NO_WINDOW',0))
        try:
            deadline = time.monotonic() + 45
            while time.monotonic() < deadline:
                if process.poll() is not None:
                    raise RuntimeError('无参数主程序提前退出：' + str(process.returncode))
                with sqlite3.connect(target.resolve().as_uri()+'?mode=ro',uri=True,timeout=1) as conn:
                    count = conn.execute('SELECT count(*) FROM cards').fetchone()[0]
                if count == 20:
                    time.sleep(1)
                    if process.poll() is not None:raise RuntimeError('内置导入后主窗口启动失败')
                    break
                time.sleep(.15)
            else:
                raise RuntimeError('无参数主程序未自动补入20张内置方法')
        finally:
            if process.poll() is None:
                process.terminate()
                process.wait(timeout=10)
        current = Store(target)
        if (len(current.list_cards()) != 20 or current.list_cards(True)
                or current.get_setting('normal_launch_preservation') != {'value':'保留正常入口的设置'}
                or current.get_review(review)['submission']['answer'] != '正常入口原答'
                or current.get_article('normal-user-article')['body'] != '原有正文保持不变。'):
            raise RuntimeError('无参数主程序补充内置方法时未保留原有数据')
    if len(list(target.parent.glob('内置文章方法导入前备份-*.db'))) != 1:
        raise RuntimeError('无参数重复启动重复导入或缺少备份')
    return {'no_command_line_arguments':True, 'uses_default_data_location':True,
            'test_localappdata':str(local), 'launches':2, 'methods':20,
            'existing_data_preserved':True, 'backup_once':True, 'cloud_calls':0}


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--test-report',type=Path,default=ROOT/'validation'/f'tests-v{__version__}.xml')
    args=parser.parse_args()
    xml=ElementTree.parse(args.test_report).getroot()
    suites=[xml] if xml.tag=='testsuite' else list(xml.iter('testsuite'))
    counts={key:sum(int(suite.get(key,0)) for suite in suites) for key in ('tests','failures','errors','skipped')}
    if not counts['tests'] or counts['failures'] or counts['errors']:
        raise RuntimeError('须先通过完整测试，再验证发布包：'+str(counts))
    counts['passed']=counts['tests']-counts['skipped']
    folder=ROOT/'dist'/'ShenlunStudio'
    exe=folder/'ShenlunStudio.exe'
    if not exe.is_file():
        raise SystemExit('请先执行 scripts/build_release.py')
    if json.loads((folder/'build-info.json').read_text(encoding='utf8'))['app']!=__version__:
        raise RuntimeError('打包版本不一致，请重新构建')
    if (folder/'_internal/assets/builtin_article_methods.json').read_bytes()!=(ROOT/'assets/builtin_article_methods.json').read_bytes():
        raise RuntimeError('打包的内置方法数据不一致，请重新构建')
    evidence=folder/'验证资料'
    evidence.mkdir(exist_ok=True)
    shutil.copy2(args.test_report,evidence/args.test_report.name)
    source_files={path.relative_to(ROOT/'shenlun'):path for path in (ROOT/'shenlun').rglob('*.py')}
    bundled_files={path.relative_to(folder/'source'/'shenlun'):path for path in (folder/'source'/'shenlun').rglob('*.py')}
    if source_files.keys()!=bundled_files.keys() or any(source_files[key].read_bytes()!=bundled_files[key].read_bytes() for key in source_files):
        raise RuntimeError('发布包源码与当前源码不一致，请重新构建')
    work=ROOT/'build-tmp'/('release-'+datetime.now().strftime('%Y%m%d-%H%M%S'))
    work.mkdir(parents=True)
    fixtures=intake_fixtures(work)
    reports=[]
    clean_env={key:value for key,value in os.environ.items()
               if key.upper() not in {'PYTHONPATH','PYTHONHOME','QT_PLUGIN_PATH','QT_QPA_PLATFORM_PLUGIN_PATH','QML2_IMPORT_PATH'}}
    windows=Path(os.environ.get('WINDIR','C:/Windows'))
    clean_env['PATH']=os.pathsep.join((str(windows/'System32'),str(windows)))
    clean_env.update(TEMP=str(work),TMP=str(work),PYTHON_KEYRING_BACKEND='keyring.backends.null.Keyring')
    for platform in ('offscreen','windows'):
        target=work/(platform+'.json')
        screenshot=work/(platform+'.png')
        command=[str(exe),'--data-dir',str(work/'data'),'--smoke-test',
                 '--screenshot',str(screenshot),'--smoke-report',str(target),
                 '--smoke-import',*[str(path) for path in fixtures]]
        subprocess.run(command,cwd=folder,env={**clean_env,'QT_QPA_PLATFORM':platform},
                       creationflags=getattr(subprocess,'CREATE_NO_WINDOW',0),timeout=110,check=True)
        report=json.loads(target.read_text(encoding='utf8'))
        if not report['ok'] or report['pages']!=6 or report['workers']!=0 or report['model_configured']:
            raise RuntimeError('打包启动状态异常：'+str(report))
        builtin=report.get('builtin_articles',{})
        if (builtin.get('cards')!=20 or builtin.get('refs')!=131 or builtin.get('enabled')!=0
                or not builtin.get('article_tab_is_default') or not builtin.get('all_evidence_valid')):
            raise RuntimeError('主程序未内置完整文章方法库：'+str(builtin))
        views=report['workbench']
        if len(views['pages'])!=6 or len(views['minimum_pages'])!=6:
            raise RuntimeError('打包工作区截图不完整')
        if any(page['width']!=1050 or page['height']!=720 for page in views['minimum_pages']):
            raise RuntimeError('工作区无法适配1050×720最小窗口')
        practice=report.get('practice',{})
        if not practice.get('ok') or not practice.get('dirty_switch_archives_latest') or practice.get('linked_attempts')!=2 or practice.get('network_model_calls')!=0:
            raise RuntimeError('打包后的练习闭环验收失败：'+str(practice))
        if practice.get('rubric_templates')!=35 or not practice.get('template_selection_recovered') or not practice.get('custom_override_protected'):
            raise RuntimeError('打包后的考试量表模板验收失败：'+str(practice))
        image_ocr=report.get('image_ocr',{})
        if not image_ocr.get('ok') or image_ocr.get('native_segments',0)<3 or image_ocr.get('loopback_image_calls')!=1 or image_ocr.get('cloud_model_calls')!=0 or image_ocr.get('assessment_writes')!=0:
            raise RuntimeError('打包后的增强OCR/当前区域图像传输验收失败：'+str(image_ocr))
        evaluation=report.get('evaluation',{})
        if not evaluation.get('ok') or not evaluation.get('synthetic_fixtures_only') or evaluation.get('loopback_calls')!=3 or evaluation.get('usage_tokens')!=60 or evaluation.get('ordinary_reviews')!=0 or evaluation.get('independent_real_human_pairs')!=0 or evaluation.get('mae') is not None or evaluation.get('cloud_calls')!=0 or not evaluation.get('new_tables_backup_restore'):
            raise RuntimeError('打包后的参考要点/冻结重复评测验收失败：'+str(evaluation))
        teaching=report.get('teaching',{})
        deletion=report.get('practice_deletion',{})
        if (deletion.get('cloud_calls')!=0 or any(not deletion.get(key) for key in
                ('ok','synthetic_fixtures_only','cancel_preserved','single_and_bulk_reviews',
                 'filtered_question_selection','single_and_bulk_questions','reports_preserved',
                 'workspace_preserved','stale_confirmation_rejected','restart_and_backup'))):
            raise RuntimeError('打包后训练记录与题库删除验收失败：'+str(deletion))
        workflow=report.get('method_workflow',{})
        if (workflow.get('loopback_calls')!=8 or workflow.get('cloud_calls')!=0 or workflow.get('human_accuracy_verified') is not False
            or any(not workflow.get(key) for key in ('ok','synthetic_fixtures_only','filters_scoped','preset_saved_from_checks',
                'exact_preview','cancel_before_upload','paired_frozen_calls','missing_humans_not_accuracy','backup_restore'))):
            raise RuntimeError('打包后方法筛选、组合和对照工作流验收失败：'+str(workflow))
        library=report.get('method_library',{})
        if (not library.get('ok') or not library.get('synthetic_fixtures_only') or library.get('cloud_calls')!=0
                or any(not library.get(key) for key in ('create_and_edit','precise_evidence','delete_cancel_preserved',
                    'article_and_teacher_delete','restored_disabled','backup_restore','historical_snapshot_preserved',
                    'manual_edits_preserved','concurrent_enable_guarded','permanent_delete',
                    'purge_cancel_preserved','purged_history_preserved'))):
            raise RuntimeError('打包后方法库新建、删除和恢复验收失败：'+str(library))
        batch=library.get('batch',{})
        if (batch.get('cloud_calls')!=0 or any(not batch.get(key) for key in ('ok','synthetic_fixtures_only',
            'checkbox_selection','filtered_select_all','cancel_preserved','article_eligibility_checked',
            'teaching_audit_saved','batch_disable','per_item_results','backup_restore'))):
            raise RuntimeError('打包后方法批量启停验收失败：'+str(batch))
        if (not teaching.get('ok') or not teaching.get('synthetic_fixtures_only') or teaching.get('catalog_drafts')!=24
                or teaching.get('community_drafts')!=12
                or teaching.get('initial_enabled')!=0 or teaching.get('loopback_calls')!=3 or teaching.get('usage_tokens')!=60
                or teaching.get('single_grade_calls')!=1 or teaching.get('compare_grade_calls')!=1
                or not teaching.get('frozen_after_disable') or not teaching.get('backup_restore')
                or teaching.get('ordinary_reviews')!=0 or teaching.get('cloud_calls')!=0):
            raise RuntimeError('打包后的教学方法调用、版本冻结或恢复验收失败：'+str(teaching))
        if report['reviews']!=0 or report['stats']['reviews']!=0:
            raise RuntimeError('合成验收记录混入了正式资料副本')
        intake=report['intake']
        if intake['pages']!=8 or intake['errors'] or not intake['cancel_preserved_editors']:
            raise RuntimeError('打包导入状态异常：'+str(intake))
        if sum('OCR' in method for method in intake['methods'])!=5 or '社区' not in intake['text'] or '电话' not in intake['text'] or '老年群众' not in intake['text']:
            raise RuntimeError('打包后的真实Word/PDF/中文OCR读取不完整')
        reports.append(report)
        store=Store(work/'data'/'studio.db')
        if platform=='offscreen':
            store.set_setting('release_preservation_probe',{'value':'保留本地数据','version':1})
        elif store.get_setting('release_preservation_probe')!={'value':'保留本地数据','version':1}:
            raise RuntimeError('重复启动覆盖了已有本地数据')
    if reports[0]['stats']!=reports[1]['stats']:
        raise RuntimeError('重复启动的语料统计不一致')
    with sqlite3.connect(work/'data'/'studio.db') as db:
        integrity=db.execute('PRAGMA quick_check').fetchone()[0]
    if integrity!='ok':
        raise RuntimeError('用户数据副本完整性异常')
    normal_launch=verify_normal_launch(exe,folder,work,clean_env)
    output=ROOT.parent/'output'/'软件'
    output.mkdir(parents=True,exist_ok=True)
    previews=folder/'界面预览'
    previews.mkdir(exist_ok=True)
    for index in range(1,7):
        shutil.copy2(work/f'windows-pages-{index}.png',previews/f'{index:02d}-工作区.png')
    shutil.copy2(work/'windows-intake.png',previews/'07-导入核对.png')
    shutil.copy2(work/'windows-ocr-region.png',previews/'08-图片选区.png')
    shutil.copy2(work/'windows-ocr-candidate.png',previews/'09-识别候选核对.png')
    shutil.copy2(work/'windows-reference-points.png',previews/'10-人工参考要点.png')
    shutil.copy2(work/'windows-evaluation.png',previews/'11-小规模评测.png')
    shutil.copy2(work/'windows-teaching-methods.png',previews/'12-名师方法库.png')
    shutil.copy2(work/'windows-teaching-report.png',previews/'13-方法应用报告.png')
    shutil.copy2(work/'windows-method-library-crud.png',previews/'14-方法管理.png')
    shutil.copy2(work/'windows-method-editor.png',previews/'15-手动方法编辑.png')
    shutil.copy2(work/'windows-batch-articles.png',previews/'16-文章方法批量启停.png')
    shutil.copy2(work/'windows-batch-teaching.png',previews/'17-教学方法批量启停.png')
    shutil.copy2(work/'windows-batch-result.png',previews/'18-批量处理结果.png')
    shutil.copy2(work/'windows-workflow-filters.png',previews/'19-方法检索筛选.png')
    shutil.copy2(work/'windows-workflow-preset-manager.png',previews/'20-常用方法组合.png')
    shutil.copy2(work/'windows-workflow-training.png',previews/'21-批改方法选择.png')
    shutil.copy2(work/'windows-workflow-comparison.png',previews/'22-方法对照评测.png')
    shutil.copy2(work/'windows-delete-records.png',previews/'23-训练记录删除.png')
    shutil.copy2(work/'windows-delete-questions.png',previews/'24-题库删除.png')
    shutil.copy2(work/'windows-pages-1.png',output/f'申论研习_{__version__}_界面预览.png')
    shutil.copy2(work/'windows-pages-3.png',output/f'申论研习_{__version__}_内置文章方法.png')
    archive=output/f'申论研习_Windows_{__version__}.zip'
    print('Creating ZIP: '+str(archive),flush=True)
    with zipfile.ZipFile(archive,'w',compression=zipfile.ZIP_DEFLATED,compresslevel=6) as bundle:
        for file in sorted(folder.rglob('*')):
            if file.is_file():
                bundle.write(file,Path('ShenlunStudio')/file.relative_to(folder))
    with zipfile.ZipFile(archive) as bundle:
        damaged=bundle.testzip()
        if damaged:
            raise RuntimeError('压缩包损坏：'+damaged)
        if 'ShenlunStudio/ShenlunStudio.exe' not in bundle.namelist():
            raise RuntimeError('压缩包缺少程序')
    with archive.open('rb') as stream:
        digest=hashlib.file_digest(stream,'sha256').hexdigest()
    archive.with_suffix('.zip.sha256').write_text(digest+'  '+archive.name+'\n',encoding='utf8')
    report={'app_version':__version__,'tests':{**counts,'junit_report':str(args.test_report.resolve())},
            'packaged_launches':reports,'existing_data_preserved':True,'sqlite_integrity':integrity,
            'normal_main_exe_launch':normal_launch,
            'paid_model_semantics_tested':False,
            'source_python_files_verified':len(source_files),
            'archive':str(archive),'bytes':archive.stat().st_size,'sha256':digest,'zip_crc_ok':True,
            'source':str(ROOT),'executable':str(exe)}
    validation=ROOT/'validation'/f'release-verification-v{__version__}.json'
    validation.write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf8')
    print(json.dumps({'verified':True,'archive':str(archive),'bytes':archive.stat().st_size,
                      'stats':reports[1]['stats'],'sha256':digest},ensure_ascii=True),flush=True)


if __name__=='__main__':
    main()
