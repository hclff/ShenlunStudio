"""Real file-format, mixed PDF and native local OCR intake checks."""
import importlib
import importlib.util
import json
from pathlib import Path
import sys
import threading
import time

import pytest
from PySide6.QtCore import QByteArray,QBuffer,QIODevice,QSize
from PySide6.QtGui import QImage,QPainter,QFont,QPdfWriter
from PySide6.QtWidgets import QApplication


def backend():
    assert importlib.util.find_spec('shenlun.intake') is not None, 'diverse local intake module is required'
    return importlib.import_module('shenlun.intake')


@pytest.fixture(scope='session')
def app():
    instance=QApplication.instance() or QApplication([])
    from shenlun.ui.widgets import ensure_fonts
    ensure_fonts()
    return instance


def image_bytes(app):
    image=QImage(1600,280,QImage.Format.Format_RGB32)
    image.fill('white')
    painter=QPainter(image)
    painter.setPen('black');painter.setFont(QFont('Microsoft YaHei',46))
    painter.drawText(50,145,'给定材料：社区保留电话服务。')
    painter.end()
    data=QByteArray();buffer=QBuffer(data);buffer.open(QIODevice.OpenModeFlag.WriteOnly)
    assert image.save(buffer,'PNG')
    buffer.close()
    return bytes(data)


@pytest.fixture
def mixed_pdf(tmp_path,app):
    path=tmp_path/'混合试卷.pdf'
    writer=QPdfWriter(str(path));writer.setResolution(100)
    painter=QPainter(writer);painter.setFont(QFont('Microsoft YaHei',36));painter.drawText(60,120,'第一材料：保留人工窗口，服务老人。')
    assert writer.newPage()
    painter.drawImage(0,80,QImage.fromData(image_bytes(app)).scaledToWidth(700))
    painter.end()
    return path


def test_gb18030_text_and_batch_order(tmp_path):
    module=backend()
    first=tmp_path/'乙.txt';first.write_bytes('群众服务。'.encode('gb18030'))
    second=tmp_path/'甲.md';second.write_text('另一则材料。',encoding='utf8')
    result=module.read_inputs([first,second])
    assert [page.text for page in result.pages]==['群众服务。','另一则材料。']
    assert not result.errors


def test_docx_preserves_paragraph_table_order(tmp_path):
    module=backend()
    from docx import Document
    document=Document();document.add_paragraph('题干在前')
    table=document.add_table(rows=1,cols=2);table.cell(0,0).text='对象';table.cell(0,1).text='群众'
    document.add_paragraph('要求在后')
    path=tmp_path/'练习.docx';document.save(path)
    result=module.read_inputs([path])
    assert not result.errors
    text=result.pages[0].text
    assert text.index('题干在前')<text.index('对象')<text.index('群众')<text.index('要求在后')


def test_structured_question_list_is_selectable_without_raw_json(tmp_path):
    module=backend()
    path=tmp_path/'题库.json'
    path.write_text(json.dumps([{'question':'概括做法','materials':'开通电话','max_score':20,'api_key':'never-copy'},
                               {'question':'解释关系','materials':'线上线下','answer':'二者互补。'}],ensure_ascii=False),encoding='utf8')
    result=module.read_inputs([path])
    assert len(result.pages)==2
    assert result.pages[0].fields['question']=='概括做法'
    assert 'api_key' not in result.pages[0].fields
    assert 'never-copy' not in result.pages[0].text
    assert '二者互补。' in result.pages[1].text


def test_unrecognized_json_stays_explicit_plain_text(tmp_path):
    module=backend()
    path=tmp_path/'data.json';path.write_text('{"title":"原文","body":"材料"}',encoding='utf8')
    result=module.read_inputs([path])
    assert not result.pages[0].fields
    assert 'body' in result.pages[0].text
    assert result.pages[0].warnings


def test_one_bad_file_preserves_other_successful_files(tmp_path):
    module=backend()
    bad=tmp_path/'旧格式.doc';bad.write_bytes(b'old format')
    good=tmp_path/'正常.txt';good.write_text('真实材料',encoding='utf8')
    result=module.read_inputs([bad,good])
    assert len(result.pages)==1 and result.pages[0].text=='真实材料'
    assert any('docx' in message for message in result.errors)


def test_file_limit_rejected_before_read(tmp_path):
    module=backend()
    paths=[tmp_path/f'{i}.txt' for i in range(21)]
    with pytest.raises(ValueError,match='20'):
        module.read_inputs(paths)


def test_oversize_file_is_reported_without_allocating_body(tmp_path):
    module=backend();path=tmp_path/'large.txt'
    with path.open('wb') as stream:
        stream.seek(26*1024*1024);stream.write(b'x')
    result=module.read_inputs([path])
    assert not result.pages and any('25' in e for e in result.errors)


def test_mixed_pdf_ocr_only_image_page(mixed_pdf,monkeypatch):
    module=backend();import shenlun.ocr as ocr
    calls=[]
    def recognize(path,cancel=None,timeout=45):
        calls.append(Path(path));return {'text':'扫描材料已识别。','language':'zh-Hans-CN'}
    monkeypatch.setattr(ocr,'recognize_png',recognize)
    result=module.read_inputs([mixed_pdf])
    assert not result.errors and len(result.pages)==2
    assert '人工窗口' in result.pages[0].text
    assert result.pages[1].text=='扫描材料已识别。'
    assert len(calls)==1
    assert all(page.preview_png for page in result.pages)
    assert result.pages[1].warnings


def test_pdf_page_spec_selects_only_requested_page(mixed_pdf,monkeypatch):
    module=backend();import shenlun.ocr as ocr
    monkeypatch.setattr(ocr,'recognize_png',lambda *a,**k:{'text':'第二页。','language':'zh-Hans-CN'})
    result=module.read_inputs([mixed_pdf],page_spec='2')
    assert len(result.pages)==1 and result.pages[0].number==2


def test_invalid_page_range_does_not_silently_truncate(mixed_pdf):
    module=backend()
    result=module.read_inputs([mixed_pdf],page_spec='1,99')
    assert not result.pages and result.errors


def test_force_ocr_can_replace_bad_existing_pdf_text(mixed_pdf,monkeypatch):
    module=backend();import shenlun.ocr as ocr
    calls=[]
    def recognize(*a,**k):
        calls.append(1);return {'text':'强制识别结果。','language':'zh-Hans-CN'}
    monkeypatch.setattr(ocr,'recognize_png',recognize)
    result=module.read_inputs([mixed_pdf],force_ocr=True)
    assert len(calls)==2 and all(p.text=='强制识别结果。' for p in result.pages)


def test_clipboard_image_keeps_original_preview(app,monkeypatch):
    module=backend();import shenlun.ocr as ocr
    monkeypatch.setattr(ocr,'recognize_png',lambda *a,**k:{'text':'社区保留电话服务。','language':'zh-Hans-CN'})
    result=module.read_image_bytes(image_bytes(app))
    assert result.pages[0].text=='社区保留电话服务。'
    assert result.pages[0].source=='剪贴板图片'
    assert result.pages[0].preview_png


def test_cancellation_never_returns_partial_success(tmp_path):
    module=backend()
    from shenlun.model import CancelledError
    event=threading.Event();event.set()
    with pytest.raises(CancelledError):
        module.read_inputs([tmp_path/'any.txt'],cancel=event)


def test_helper_cancel_kills_child(tmp_path,monkeypatch):
    backend();import shenlun.ocr as ocr
    from shenlun.model import CancelledError
    event=threading.Event();timer=threading.Timer(.25,event.set)
    actual=ocr.subprocess.Popen;children=[]
    def spawning(*args,**kwargs):
        process=actual(*args,**kwargs);children.append(process);return process
    monkeypatch.setattr(ocr.subprocess,'Popen',spawning)
    timer.start()
    try:
        with pytest.raises(CancelledError):
            ocr.run_hidden_json([sys.executable,'-c','import sys,time; sys.stdin.read(); time.sleep(30)'],{},cancel=event,timeout=10)
    finally:
        timer.cancel()
    assert children and children[0].poll() is not None


@pytest.mark.skipif(sys.platform!='win32',reason='Native Windows Chinese OCR integration')
def test_actual_native_chinese_ocr(app,tmp_path):
    backend();import shenlun.ocr as ocr
    path=tmp_path/"材料'a&b.png";path.write_bytes(image_bytes(app))
    result=ocr.recognize_png(path)
    assert '社区' in result['text'] and '电话' in result['text']
    assert result['language'].startswith('zh')


@pytest.mark.skipif(sys.platform!='win32',reason='Native Windows scan PDF OCR integration')
def test_actual_scanned_pdf_ocr(mixed_pdf):
    result=backend().read_inputs([mixed_pdf],page_spec='2')
    assert not result.errors and len(result.pages)==1
    assert '社区' in result.pages[0].text and '电话' in result.pages[0].text
    assert 'OCR' in result.pages[0].method and result.pages[0].preview_png


def test_helper_timeout_kills_child(monkeypatch):
    backend();import shenlun.ocr as ocr
    actual=ocr.subprocess.Popen;children=[]
    def spawning(*args,**kwargs):
        process=actual(*args,**kwargs);children.append(process);return process
    monkeypatch.setattr(ocr.subprocess,'Popen',spawning)
    with pytest.raises(ValueError,match='超时'):
        ocr.run_hidden_json([sys.executable,'-c','import sys,time; sys.stdin.read(); time.sleep(30)'],{},timeout=.3)
    assert children and children[0].poll() is not None


@pytest.mark.parametrize('suffix,format_name',[('.webp','WEBP'),('.tif','TIFF'),('.tiff','TIFF')])
def test_webp_and_multipage_tiff_decode_without_optional_qt_plugins(app,tmp_path,monkeypatch,suffix,format_name):
    from PIL import Image
    module=backend();import shenlun.ocr as ocr
    path=tmp_path/('材料'+suffix)
    first=Image.new('RGB',(180,90),'white');second=Image.new('RGB',(180,90),'yellow')
    if format_name=='TIFF':first.save(path,format_name,save_all=True,append_images=[second])
    else:first.save(path,format_name)
    calls=[]
    def recognize(*a,**k):
        calls.append(1);return {'text':f'材料{len(calls)}','language':'zh-Hans-CN'}
    monkeypatch.setattr(ocr,'recognize_png',recognize)
    result=module.read_inputs([path])
    assert not result.errors
    assert len(result.pages)==(2 if format_name=='TIFF' else 1)
    assert [p.number for p in result.pages]==list(range(1,len(result.pages)+1))
    assert all(p.preview_png for p in result.pages)


def test_batch_limit_stops_ocr_before_extra_pdf_pages(app,tmp_path,monkeypatch):
    from pypdf import PdfWriter
    module=backend();import shenlun.ocr as ocr
    paths=[]
    for i in range(3):
        writer=PdfWriter()
        for _ in range(30):writer.add_blank_page(width=72,height=72)
        path=tmp_path/f'{i}.pdf';writer.write(path);paths.append(path)
    calls=[]
    def recognize(*a,**k):
        calls.append(1);return {'text':'已识别','language':'zh-Hans-CN'}
    monkeypatch.setattr(ocr,'recognize_png',recognize)
    result=module.read_inputs(paths)
    assert len(result.pages)==80 and len(calls)==80
    assert any('80' in error for error in result.errors)


def test_bad_json_overflow_record_preserves_valid_siblings(tmp_path):
    path=tmp_path/'题库.json'
    path.write_text(json.dumps([{'question':'有效题一'},{'question':'错误题','max_score':10**400},{'question':'有效题二'}],ensure_ascii=False),encoding='utf8')
    result=backend().read_inputs([path])
    assert [p.fields['question'] for p in result.pages]==['有效题一','有效题二']
    assert result.errors


def test_structured_rubric_only_keeps_supported_fields(tmp_path):
    path=tmp_path/'量表.json'
    path.write_text(json.dumps({'question':'概括做法','rubric':[{'id':'coverage','name':'要点','weight':100,'description':'完整准确','api_key':'must-not-copy'}]},ensure_ascii=False),encoding='utf8')
    result=backend().read_inputs([path])
    assert not result.errors
    assert 'api_key' not in result.pages[0].fields['rubric'][0]


@pytest.mark.parametrize('fail_every',[1,2])
def test_batch_work_limit_includes_failed_ocr_attempts(app,tmp_path,monkeypatch,fail_every):
    from pypdf import PdfWriter
    module=backend();import shenlun.ocr as ocr
    paths=[]
    for i in range(3):
        writer=PdfWriter()
        for _ in range(30):writer.add_blank_page(width=72,height=72)
        path=tmp_path/f'{i}.pdf';writer.write(path);paths.append(path)
    calls=[]
    def recognize(*a,**k):
        calls.append(1)
        if len(calls)%fail_every==0:raise ValueError('OCR unavailable')
        return {'text':'成功页','language':'zh-Hans-CN'}
    monkeypatch.setattr(ocr,'recognize_png',recognize)
    result=module.read_inputs(paths)
    assert len(calls)==80
    assert len(result.pages)==(0 if fail_every==1 else 40)
    assert result.errors
