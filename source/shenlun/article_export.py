"""Export selected articles to Markdown, JSON or Word without touching the UI."""
import json
from pathlib import Path

MARKDOWN_SUFFIXES = ('.md', '.markdown')
JSON_SUFFIXES = ('.json',)
DOCX_SUFFIXES = ('.docx',)
META_KEYS = ('id', 'title', 'author', 'journal', 'publication', 'kind', 'url')


def paragraphs_of(article):
    """Stored paragraphs first, then the body as a fallback."""
    result = [str(item.get('text', '')).strip() for item in article.get('paragraphs') or []
              if str(item.get('text', '')).strip()]
    if result:
        return result
    return [line.strip() for line in str(article.get('body') or '').split('\n') if line.strip()]


def _front_matter(article):
    lines = ['---']
    for key in META_KEYS:
        lines.append(f'{key}: {json.dumps(str(article.get(key) or ""), ensure_ascii=False)}')
    tags = '[' + ', '.join(json.dumps(str(tag), ensure_ascii=False)
                           for tag in article.get('tags') or []) + ']'
    lines.append(f'tags: {tags}')
    lines.append('---')
    return '\n'.join(lines)


def export_markdown(articles, path):
    """One Markdown file; several articles are separated by a horizontal rule."""
    blocks = []
    for article in articles:
        body = '\n\n'.join(paragraphs_of(article))
        blocks.append(_front_matter(article) + '\n\n' + body)
    Path(path).write_text('\n\n---\n\n'.join(blocks) + '\n', encoding='utf-8')
    return str(path)


def export_json(articles, path):
    """The stored ``data`` payloads verbatim, one array per file."""
    Path(path).write_text(json.dumps(list(articles), ensure_ascii=False, indent=2), encoding='utf-8')
    return str(path)


def export_docx(articles, path):
    """Word output: heading, metadata line, tags and every paragraph."""
    from docx import Document

    document = Document()
    for index, article in enumerate(articles):
        if index:
            document.add_page_break()
        document.add_heading(str(article.get('title') or '未命名文章'), level=1)
        meta = ' · '.join(part for part in (str(article.get('author') or ''),
                                            str(article.get('journal') or ''),
                                            str(article.get('publication') or '')[:10],
                                            str(article.get('kind') or '')) if part)
        if meta:
            document.add_paragraph(meta)
        if article.get('tags'):
            document.add_paragraph('标签：' + '、'.join(article['tags']))
        for text in paragraphs_of(article):
            document.add_paragraph(text)
    document.save(str(path))
    return str(path)


def export_articles(articles, path):
    """Dispatch on the file suffix; the caller reports the returned path."""
    articles = list(articles)
    if not articles:
        raise ValueError('请先选择要导出的文章。')
    suffix = Path(path).suffix.lower()
    if suffix in MARKDOWN_SUFFIXES:
        return export_markdown(articles, path)
    if suffix in JSON_SUFFIXES:
        return export_json(articles, path)
    if suffix in DOCX_SUFFIXES:
        return export_docx(articles, path)
    raise ValueError('仅支持导出为 Markdown、JSON 或 Word（.docx）')
