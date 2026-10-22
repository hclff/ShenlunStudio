"""Read local documents into plain paragraphs without touching the database.

Every reader returns the same shape so the importers can stay format agnostic:

    {'title': str, 'author': str, 'publication': str, 'tags': list,
     'body': str, 'paragraphs': [str, ...], 'format': str}

Paragraph ids are assigned later by ``normalize_article`` (``<article_id>:pN``),
which keeps new formats valid evidence for existing method cards.
"""
import json
import posixpath
import re
import zipfile
from datetime import datetime
from pathlib import Path
from urllib.parse import unquote
from xml.etree import ElementTree as ET

MAX_FILE_BYTES = 25 * 1024 * 1024
DOCUMENT_SUFFIXES = ('.txt', '.md', '.pdf', '.docx', '.epub', '.json')
READABLE_SUFFIXES = ('.txt', '.md', '.pdf', '.docx', '.epub')
FLAT_SUFFIXES = ('.txt', '.md')

_XHTML_TYPES = ('application/xhtml+xml', 'text/html', 'application/xml')
_BLOCK_TAGS = ('p', 'h1', 'h2', 'h3', 'h4', 'h5', 'h6', 'li', 'blockquote', 'pre', 'td')
_SPACE = re.compile(r'\s+')
_TAG_SPLIT = re.compile(r'[,，、;；]')


def clean_text(text):
    """Collapse whitespace while keeping CJK characters untouched."""
    return _SPACE.sub(' ', str(text or '')).strip()


def read_document(path) -> dict:
    """Parse one local document; JSON is handled by ``importers``."""
    path = Path(path)
    if not path.is_file():
        raise ValueError(f'找不到文件：{path.name}')
    if path.stat().st_size > MAX_FILE_BYTES:
        raise ValueError('文件超过25MB限制；官方来源清单请使用专用导入')
    suffix = path.suffix.lower()
    if suffix in FLAT_SUFFIXES:
        return _read_flat(path)
    if suffix == '.pdf':
        return _read_pdf(path)
    if suffix == '.docx':
        return _read_docx(path)
    if suffix == '.epub':
        return _read_epub(path)
    if suffix == '.json':
        raise ValueError('JSON 请使用“已有清单/文章 JSON”导入')
    raise ValueError('仅支持 TXT、MD、PDF、DOCX、EPUB')


def read_text(path) -> str:
    """Whole-document text, used where the caller only wants the raw string."""
    return read_document(path)['body']


def _record(paragraphs, *, title, author='', publication='', tags=(), fmt=''):
    body = '\n'.join(paragraphs)
    return {'title': clean_text(title) or '未命名文章', 'author': clean_text(author),
            'publication': _stamp(publication), 'tags': list(tags), 'body': body,
            'paragraphs': paragraphs, 'format': fmt}


def _stamp(value):
    """Normalise a date to the project's ``YYYY-MM-DD HH:MM:SS`` form."""
    text = str(value or '').strip()
    match = re.search(r'(\d{4})[-/年.]?(\d{1,2})[-/月.]?(\d{1,2})', text)
    if not match:
        return ''
    year, month, day = (int(part) for part in match.groups())
    try:
        datetime(year, month, day)
    except ValueError:
        return ''
    return f'{year:04d}-{month:02d}-{day:02d} 00:00:00'


def _paragraphs(text):
    return [line.strip() for line in str(text or '').replace('\r\n', '\n').replace('\r', '\n').split('\n')
            if line.strip()]


def _read_flat(path):
    text = path.read_text(encoding='utf-8-sig')
    title, author, publication, tags, body = path.stem, '', '', [], text
    if path.suffix.lower() == '.md':
        title, author, publication, tags, body = _markdown_parts(text, path.stem)
    paragraphs = _paragraphs(body)
    if not paragraphs:
        raise ValueError('文件没有正文')
    return _record(paragraphs, title=title, author=author, publication=publication,
                   tags=tags, fmt=path.suffix.lower().lstrip('.'))


def _markdown_parts(text, stem):
    """Front matter wins, then the first level-one heading, then the file name."""
    title, author, publication, tags, body = stem, '', '', [], text
    if text.startswith('---'):
        end = text.find('\n---', 3)
        if end > 0:
            head, body = text[3:end], text[end + 4:]
            for line in head.split('\n'):
                if ':' not in line:
                    continue
                key, value = (part.strip() for part in line.split(':', 1))
                value = value.strip('"\' ')
                if key in ('title', 'name'):
                    title = value or title
                elif key in ('author', 'creator'):
                    author = value
                elif key in ('date', 'publication'):
                    publication = value
                elif key in ('tags', 'keywords'):
                    tags = [item.strip() for item in _TAG_SPLIT.split(value.strip('[]'))
                            if item.strip()]
    if not title or title == stem:
        heading = re.search(r'^#\s+(.+)$', body, re.MULTILINE)
        if heading:
            title = heading.group(1).strip()
            body = body[:heading.start()] + body[heading.end():]
    return title, author, publication, tags, body


def _read_pdf(path):
    from pypdf import PdfReader

    try:
        reader = PdfReader(path)
        text = '\n'.join(page.extract_text() or '' for page in reader.pages)
    except Exception as exc:  # noqa: BLE001 - surface a readable reason
        raise ValueError(f'PDF 无法解析：{exc}') from exc
    if not text.strip():
        raise ValueError('PDF没有可提取文字，扫描件需先使用外部OCR')
    paragraphs = _paragraphs(text)
    title = path.stem
    try:
        title = (reader.metadata or {}).get('/Title') or path.stem
    except Exception:  # noqa: BLE001 - metadata is optional
        pass
    return _record(paragraphs, title=title, fmt='pdf')


def _read_docx(path):
    from docx import Document

    try:
        document = Document(str(path))
    except Exception as exc:  # noqa: BLE001 - python-docx raises many subclasses
        raise ValueError(f'DOCX 无法解析：{exc}') from exc
    props = document.core_properties
    paragraphs = []
    heading = ''
    for paragraph in document.paragraphs:
        text = clean_text(paragraph.text)
        if not text:
            continue
        style = str(getattr(paragraph.style, 'name', '') or '')
        if not heading and style.lower().startswith(('heading 1', '标题 1', 'title')):
            heading = text
            continue
        paragraphs.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [clean_text(cell.text) for cell in row.cells]
            line = ' '.join(cell for cell in cells if cell)
            if line:
                paragraphs.append(line)
    if not paragraphs:
        raise ValueError('DOCX 没有可提取正文')
    return _record(paragraphs, title=props.title or heading or path.stem,
                   author=props.author or '', publication=props.created or props.modified or '',
                   tags=[item.strip() for item in str(props.keywords or '').split(',') if item.strip()],
                   fmt='docx')


def _epub_package(archive):
    """Locate the OPF package and return its path plus parsed root element."""
    container = ET.fromstring(archive.read('META-INF/container.xml'))
    for element in container.iter():
        if element.tag.rsplit('}', 1)[-1] == 'rootfile':
            opf_path = element.get('full-path')
            if opf_path:
                return opf_path, ET.fromstring(archive.read(opf_path))
    raise ValueError('EPUB 缺少 OPF 目录，无法解析')


def _epub_metadata(package):
    """Pull dc:title / dc:creator / dc:date out of the OPF metadata block."""
    meta = {'title': '', 'author': '', 'date': ''}
    for element in package.iter():
        name = element.tag.rsplit('}', 1)[-1]
        text = clean_text(element.text)
        if name == 'title' and text and not meta['title']:
            meta['title'] = text
        elif name == 'creator' and text and not meta['author']:
            meta['author'] = text
        elif name == 'date' and text and not meta['date']:
            meta['date'] = text
    return meta


def _epub_spine(archive, opf_path, package):
    """Resolve spine documents in reading order, honouring the manifest hrefs."""
    base = posixpath.dirname(opf_path)
    manifest, order = {}, []
    for element in package.iter():
        name = element.tag.rsplit('}', 1)[-1]
        if name == 'item' and element.get('id') and element.get('href'):
            media = (element.get('media-type') or '').lower()
            properties = (element.get('properties') or '').lower()
            if media in _XHTML_TYPES and 'nav' not in properties.split():
                manifest[element.get('id')] = element.get('href')
        elif name == 'itemref' and element.get('idref'):
            order.append(element.get('idref'))
    documents = [manifest[item] for item in order if item in manifest] or list(manifest.values())
    result = []
    for href in documents:
        target = unquote(href.split('#', 1)[0])
        name = posixpath.normpath(posixpath.join(base, target))
        if name in archive.namelist():
            result.append(name)
    return result


def _read_epub(path):
    try:
        with zipfile.ZipFile(path) as archive:
            opf_path, package = _epub_package(archive)
            meta = _epub_metadata(package)
            documents = _epub_spine(archive, opf_path, package)
            paragraphs, seen = [], set()
            for name in documents:
                try:
                    root = ET.fromstring(archive.read(name))
                except ET.ParseError:
                    continue
                for element in root.iter():
                    if element.tag.rsplit('}', 1)[-1] not in _BLOCK_TAGS:
                        continue
                    text = clean_text(''.join(element.itertext()))
                    if text and text not in seen:
                        seen.add(text)
                        paragraphs.append(text)
    except zipfile.BadZipFile as exc:
        raise ValueError('EPUB 文件已损坏或不是 ZIP 容器') from exc
    except KeyError as exc:
        raise ValueError(f'EPUB 结构不完整：缺少 {exc}') from exc
    if not paragraphs:
        raise ValueError('EPUB 没有可提取正文')
    return _record(paragraphs, title=meta['title'] or path.stem, author=meta['author'],
                   publication=meta['date'], fmt='epub')


def read_json_articles(path):
    """JSON payloads may be one article, a list, or an ``{articles: [...]}`` wrapper."""
    path = Path(path)
    if not path.is_file():
        raise ValueError(f'找不到文件：{path.name}')
    if path.stat().st_size > MAX_FILE_BYTES:
        raise ValueError('文件超过25MB限制；官方来源清单请使用专用导入')
    try:
        data = json.loads(path.read_text(encoding='utf-8-sig'))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise ValueError(f'JSON 无法解析：{exc}') from exc
    if isinstance(data, dict) and 'issues' in data:
        raise ValueError('官方来源清单请使用专用导入')
    articles = data if isinstance(data, list) else data.get('articles', [data]) if isinstance(data, dict) else []
    if not articles or any(not isinstance(article, dict) for article in articles):
        raise ValueError('JSON必须包含文章对象或文章列表')
    return articles
